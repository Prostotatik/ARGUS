from perturb import *
import perturb, random, io
perturb.PAIRS = perturb.PAIRS[:40]
import docx, openpyxl, pymupdf
def kv(text):
    out=[]
    for l in text.split("\n"):
        m=re.match(r"^([A-Za-z][^:]{0,60}):\s?(.*)$", l)
        if m: out.append((m.group(1), m.group(2)))
    return out
title=lambda t: t.split("\n")[0]
def mk_docx(text, style):
    d=docx.Document(); d.add_paragraph(title(text))
    if style=="para":
        for l in text.split("\n")[2:]:
            d.add_paragraph(l)
    elif style=="table2":
        rows=kv(text); t=d.add_table(rows=len(rows), cols=2)
        for i,(a,b) in enumerate(rows): t.cell(i,0).text=a; t.cell(i,1).text=b
    elif style=="table4":
        rows=kv(text); t=d.add_table(rows=(len(rows)+1)//2, cols=4)
        for i,(a,b) in enumerate(rows): r,c=divmod(i,2); t.cell(r,c*2).text=a; t.cell(r,c*2+1).text=b
    elif style=="table_hdr_col":  # header row of labels, value row below (transposed)
        rows=kv(text); t=d.add_table(rows=2, cols=len(rows))
        for i,(a,b) in enumerate(rows): t.cell(0,i).text=a; t.cell(1,i).text=b
    bio=io.BytesIO(); d.save(bio); return bio.getvalue()
def mk_xlsx(text, style):
    wb=openpyxl.Workbook(); ws=wb.active; ws.append([title(text)])
    rows=kv(text)
    if style=="ab":
        for a,b in rows: ws.append([a,b])
    elif style=="bc":
        ws.append([]);
        for a,b in rows: ws.append([None,a,b])
    elif style=="hdr":
        ws.append([a for a,b in rows]); ws.append([b for a,b in rows])
    elif style=="numeric":
        for a,b in rows:
            m=re.match(r"^([\d,]+) KG$", b)
            ws.append([a, int(m.group(1).replace(",","")) if m else b])
    bio=io.BytesIO(); wb.save(bio); return bio.getvalue()
def mk_pdf(text, style):
    doc=pymupdf.open(); page=doc.new_page(); y=60
    rows=kv(text)
    if style!="scan": page.insert_text((50,y), title(text), fontsize=14); y+=30
    if style=="lines":
        for a,b in rows:
            page.insert_text((50,y), f"{a}: {b}", fontsize=9); y+=14
    elif style=="cols":
        for a,b in rows:
            page.insert_text((50,y), a, fontsize=9); page.insert_text((250,y), b, fontsize=9); y+=14
    elif style=="scan":
        pix=None
        p2=pymupdf.open(); pg=p2.new_page(); yy=60
        pg.insert_text((50,yy), title(text), fontsize=14); yy+=30
        for a,b in rows: pg.insert_text((50,yy), f"{a}: {b}", fontsize=9); yy+=14
        pix=pg.get_pixmap(dpi=100)
        page.insert_image(page.rect, pixmap=pix)
    return doc.tobytes()
def build_fmt(name, siext, blext, sistyle, blstyle):
    d = SCR/name
    if d.exists(): shutil.rmtree(d)
    (d/"inbox").mkdir(parents=True); (d/"attachments").mkdir()
    ids=[]
    for e in PAIRS:
        si=(BUNDLE/e["attachments"][0]).read_text(encoding="utf-8"); bl=(BUNDLE/e["attachments"][1]).read_text(encoding="utf-8")
        e2=dict(e); atts=[]
        for role,t,ext,st in (("SI",si,siext,sistyle),("BL",bl,blext,blstyle)):
            p=f"attachments/{e['email_id']}_{role}.{ext}"; atts.append(p)
            if ext=="docx": data=mk_docx(t,st)
            elif ext=="xlsx": data=mk_xlsx(t,st)
            elif ext=="pdf": data=mk_pdf(t,st)
            else: data=t.encode()
            (d/p).write_bytes(data)
        e2["attachments"]=atts
        (d/"inbox"/f"{e['email_id']}.json").write_text(json.dumps(e2))
        ids.append(e["email_id"])
    return d, ids
def ev2(name, *a):
    d,ids=build_fmt(name,*a)
    res=asyncio.run(run(d)); sub=build_submission(res)
    ch=[i for i in ids if (base[i]["status"], sorted(base[i]["defect_fields"]))!=(sub[i]["status"], sorted(sub[i]["defect_fields"]))]
    kinds={}
    for i in ch:
        k=(base[i]["status"],"->",sub[i]["status"],sub[i]["review_reason"]); kinds[k]=kinds.get(k,0)+1
    print(f"[{name}] changed {len(ch)}/{len(ids)} {kinds}")
    if ch: print("    e.g.", ch[0], res[ch[0]]["headline"][:160])

for nm,a in [("mixed_docx_xlsx",("docx","xlsx","table2","ab")),("scan",("pdf","pdf","scan","scan")),("scan_si_only",("pdf","txt","scan","x"))]:
    ev2(nm,*a)
